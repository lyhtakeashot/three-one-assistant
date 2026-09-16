# -*- coding: utf-8 -*-
"""按「三位一体数据采集需求表.md」§2 生成 Excel 数据采集表。
数据源：
  - live-app/open-data/schools.json（现有 15 所院校真实数据，用于预填表A/B/C/D）
  - 浙江省教育考试院 2026 三位一体名单（32 所省内地方属）+ 2026 高水平三位一体 6 所
    + 在浙综合评价招生 8 所 + 2025 省内地方属中 2026 未招的 7 所（历史备查）
"""
import json
import os

try:
    import openpyxl  # noqa
except ImportError:
    import subprocess
    import sys

    subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "openpyxl>=3.1.0"])

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BASE = r"D:\Desktop\VibeCoding\三位一体辅助系统"
SRC = os.path.join(BASE, "live-app", "open-data", "schools.json")
OUT = os.path.join(BASE, "三位一体数据采集表.xlsx")


def xl_color(css_hex: str) -> str:
    value = css_hex.removeprefix("#").upper()
    if len(value) != 6:
        raise ValueError(f"Expected #RRGGBB, got: {css_hex}")
    return "FF" + value


XL_PRIMARY = xl_color("#4472C4")
XL_TITLE_BG = xl_color("#2F5597")
XL_LIGHT = xl_color("#D9E2F3")
XL_BORDER = xl_color("#BFBFBF")
XL_WHITE = xl_color("#FFFFFF")

thin = Side(style="thin", color=XL_BORDER)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

FONT = "微软雅黑"
font_body = Font(name=FONT, size=10)
font_hdr = Font(name=FONT, size=10, bold=True, color=XL_WHITE)
font_title = Font(name=FONT, size=13, bold=True, color=XL_WHITE)

A_HDR = Alignment(horizontal="center", vertical="center", wrap_text=True)
A_LEFT = Alignment(horizontal="left", vertical="center")
A_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)
A_CENTER = Alignment(horizontal="center", vertical="center")
A_RIGHT = Alignment(horizontal="right", vertical="center")

FILL_HDR = PatternFill("solid", fgColor=XL_PRIMARY)
FILL_TITLE = PatternFill("solid", fgColor=XL_TITLE_BG)
FILL_LIGHT = PatternFill("solid", fgColor=XL_LIGHT)

# ---------- 源数据 ----------
with open(SRC, "r", encoding="utf-8") as fh:
    DATA = json.load(fh)
SCHOOLS = DATA["schools"]
COLLECTED = {s["id"] for s in SCHOOLS}

TYPE_MAP = {"ministry": "部属", "provincial": "省属"}
INTERVIEW_MAP = {"individual": "个面", "group": "群面", "both": "两者"}

# ---------- 2026 院校全名单（官方口径） ----------
CAT_DIFANG = "省内地方属三位一体"
CAT_GSP = "高水平三位一体"
CAT_ZH = "综合评价（在浙）"
CAT_LS = "历史（2026未招）"

DIFANG_2026 = [
    ("浙江工业大学", "zjut"), ("浙江师范大学", "zjnu"), ("宁波大学", "nbu"), ("杭州电子科技大学", "hdu"),
    ("浙江工商大学", "zjgsu"), ("浙江理工大学", "zstu"), ("温州医科大学", "wmu"), ("浙江海洋大学", "zjou"),
    ("浙江农林大学", "zafu"), ("浙江中医药大学", "zcmu"), ("中国计量大学", "cjlu"), ("浙江万里学院", "zwu"),
    ("浙江科技大学", "zust"), ("浙江财经大学", "zufe"), ("嘉兴大学", "zjxu"), ("浙大城市学院", "zucc"),
    ("浙大宁波理工学院", "nit"), ("杭州师范大学", "hznu"), ("湖州师范大学", "hzsf"), ("绍兴大学", "sxu"),
    ("温州大学", "wzu"), ("浙江外国语学院", "zisu"), ("浙江水利水电学院", "zjweu"), ("浙江警察学院", "zjpc"),
    ("杭州医学院", "hzmc"), ("丽水学院", "lsu"), ("嘉兴南湖学院", "jxnh"), ("温州肯恩大学", "wku"),
    ("宁波诺丁汉大学", "unnc"), ("浙江越秀外国语学院", "zyu"), ("宁波财经学院", "nbfe"), ("温州商学院", "wzbc"),
]
GSP_2026 = [
    ("复旦大学", "fudan"), ("上海交通大学", "sjtu"), ("浙江大学", "zju"),
    ("中国科学院大学", "ucas"), ("西湖大学", "westlake"), ("宁波东方理工大学", "eit"),
]
ZH_2026 = [
    ("上海纽约大学", "nyush"), ("昆山杜克大学", "dku"), ("北京外国语大学", "bfsu"),
    ("上海科技大学", "shtech"), ("南方科技大学", "sustc"), ("华南理工大学", "scut"),
    ("香港中文大学（深圳）", "cuhksz"), ("深圳北理莫斯科大学", "smbu"),
]
LS_2026 = [
    ("台州学院", "tzc"), ("衢州学院", "qzc"), ("宁波工程学院", "nbut"), ("温州理工学院", "wzut"),
    ("金华职业技术大学", "jhc"), ("宁波幼儿师范高等专科学校", "nbyz"), ("湖州学院", "hzxy"),
]


def yesno(val):
    if val is None:
        return None
    return "是" if val else "否"


def joinv(val):
    if not val:
        return None
    return ";".join(str(x) for x in val)


# ---------- 通用建表 ----------
def build_sheet(wb, name, title, headers, rows, widths, num_formats=None,
                center_cols=(), right_cols=(), wrap_cols=(), dv_list=(),
                freeze="B3", reserve=0, use_filter=True):
    ws = wb.create_sheet(name)
    ncol = len(headers)

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    tc = ws.cell(row=1, column=1, value=title)
    tc.font = font_title
    tc.fill = FILL_TITLE
    tc.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28

    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=2, column=j, value=h)
        c.font = font_hdr
        c.fill = FILL_HDR
        c.alignment = A_HDR
        c.border = BORDER
    ws.row_dimensions[2].height = 32

    n_data = len(rows)
    last_row = 2 + n_data + reserve
    for i in range(n_data + reserve):
        r = 3 + i
        rv = rows[i] if i < n_data else []
        for j in range(1, ncol + 1):
            v = rv[j - 1] if j - 1 < len(rv) else None
            c = ws.cell(row=r, column=j, value=v)
            c.font = font_body
            if j in center_cols:
                c.alignment = A_CENTER
            elif j in right_cols:
                c.alignment = A_RIGHT
            elif j in wrap_cols:
                c.alignment = A_WRAP
            else:
                c.alignment = A_LEFT
            c.border = BORDER
            if num_formats and num_formats[j - 1]:
                c.number_format = num_formats[j - 1]

    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w

    for col_idx, opts in dv_list:
        dv = DataValidation(type="list", formula1='"%s"' % opts, allow_blank=True)
        dv.error = "请从下拉列表中选择"
        dv.promptTitle = "可选值"
        dv.prompt = opts
        ws.add_data_validation(dv)
        col = get_column_letter(col_idx)
        dv.add("%s3:%s%d" % (col, col, last_row))

    if freeze:
        ws.freeze_panes = freeze
    if use_filter:
        ws.auto_filter.ref = "A2:%s%d" % (get_column_letter(ncol), last_row)

    return ws


wb = Workbook()
wb.remove(wb.active)

# =========================================================
# 表A：院校基础信息（每校 1 行）
# =========================================================
A_HEADERS = [
    "院校全称", "简称", "别名", "类型", "id",
    "校区1名称", "校区1地址", "校区2名称", "校区2地址", "官网", "招生办电话", "咨询QQ",
    "普通学费", "中外合作学费", "体检限制",
    "招生简章链接", "章程年份", "章程链接",
    "学考A", "学考B", "学考C", "学考D", "学考满分", "校测满分", "高考满分",
    "学考权重", "校测权重", "高考权重",
    "有笔试", "有面试", "有体测", "笔试科目", "面试形式", "校测内容摘要", "备考建议",
    "转专业受限", "转专业说明",
    "综合满意度", "环境满意度", "生活满意度", "满意度来源",
    "宿舍描述", "宿舍评分", "宿舍亮点", "宿舍不足",
]
A_WIDTHS = [
    14, 8, 14, 8, 8,
    12, 16, 12, 16, 26, 14, 12,
    16, 18, 22,
    26, 10, 26,
    7, 7, 7, 7, 9, 9, 9,
    9, 9, 9,
    8, 8, 8, 16, 10, 40, 34,
    10, 34,
    10, 10, 10, 12,
    40, 10, 24, 24,
]
A_NUMFMT = [
    None, None, None, None, None,
    None, None, None, None, None, None, None,
    None, None, None,
    None, "0", None,
    "0", "0", "0", "0", "0", "0", "0",
    "0.00", "0.00", "0.00",
    None, None, None, None, None, None, None,
    None, None,
    "0.0", "0.0", "0.0", None,
    None, "0.0", None, None,
]
A_ROWS = []
for s in SCHOOLS:
    info = s.get("info", {})
    camps = info.get("campuses", [])
    c1 = camps[0] if len(camps) > 0 else {}
    c2 = camps[1] if len(camps) > 1 else {}
    f = s.get("formula", {})
    xk = f.get("xuekao", {})
    w = f.get("weights", {})
    ef = s.get("examFormat", {})
    tr = s.get("transferRestriction", {})
    sat = s.get("satisfaction", {})
    dor = s.get("dormitory", {})
    fs = s.get("formulaSource", {}) or {}
    A_ROWS.append([
        s.get("name"), s.get("shortName"), joinv(s.get("aliases")),
        TYPE_MAP.get(s.get("type"), s.get("type")), s.get("id"),
        c1.get("name"), c1.get("address"), c2.get("name"), c2.get("address"),
        info.get("website"), info.get("admissionsPhone"), info.get("consultQQ"),
        info.get("tuitionGeneral"), info.get("tuitionSinoForeign"), info.get("healthRestrictions"),
        s.get("brochureUrl"), fs.get("year"), fs.get("url"),
        xk.get("A"), xk.get("B"), xk.get("C"), xk.get("D"), xk.get("fullScore"),
        f.get("xiaokao", {}).get("fullScore"), f.get("gaokao", {}).get("fullScore"),
        w.get("xuekao"), w.get("xiaokao"), w.get("gaokao"),
        yesno(ef.get("hasWrittenTest")), yesno(ef.get("hasInterview")), yesno(ef.get("hasPhysicalTest")),
        joinv(ef.get("writtenTestSubjects")),
        INTERVIEW_MAP.get(ef.get("interviewFormat"), ef.get("interviewFormat")),
        ef.get("contentSummary"), ef.get("tips"),
        yesno(tr.get("restricted")), tr.get("detail"),
        sat.get("overall"), sat.get("environment"), sat.get("life"), sat.get("source"),
        dor.get("description"), dor.get("score"), joinv(dor.get("highlights")), joinv(dor.get("drawbacks")),
    ])

# 2026 全部招生院校 = 32 + 6 + 8 = 46，已填 15 → 预留 31
build_sheet(
    wb, "表A-院校基础信息",
    "表A：院校基础信息（每校 1 行）· 已填 15 所 + 预留 31 行（2026 共 46 所招生院校）",
    A_HEADERS, A_ROWS, A_WIDTHS, num_formats=A_NUMFMT,
    center_cols={4, 17, 19, 20, 21, 22, 23, 24, 25, 26, 27, 28, 29, 30, 31, 33, 36, 38, 39, 40, 43},
    wrap_cols={7, 9, 12, 15, 18, 32, 34, 35, 37, 42, 44, 45},
    dv_list=[(4, "省属,部属"), (29, "是,否"), (30, "是,否"), (31, "是,否"),
             (33, "个面,群面,两者"), (36, "是,否")],
    freeze="B3", reserve=31,
)

# =========================================================
# 表B：招生专业（每专业 1 行）
# =========================================================
B_HEADERS = ["院校全称", "专业名称", "专业类别", "选科要求", "招生计划数", "最低分年份", "最低综合分", "来源链接"]
B_WIDTHS = [16, 22, 12, 14, 12, 12, 12, 30]
B_NUMFMT = [None, None, None, None, "0", "0", "0.0", None]
B_ROWS = []
for s in SCHOOLS:
    for m in s.get("majors", []):
        st = m.get("admissionStats", {}) or {}
        B_ROWS.append([
            s.get("name"), m.get("name"), m.get("category"), joinv(m.get("requiredSubjects")),
            m.get("planCount"), st.get("year"), st.get("minScore"), st.get("sourceUrl"),
        ])

build_sheet(
    wb, "表B-招生专业", "表B：招生专业（每专业 1 行）",
    B_HEADERS, B_ROWS, B_WIDTHS, num_formats=B_NUMFMT,
    center_cols={5, 6, 7}, wrap_cols={8}, freeze="B3",
)

# =========================================================
# 表C：历年竞争比（每年 1 行）
# =========================================================
C_HEADERS = ["院校全称", "年份", "报名人数", "入围人数", "录取人数", "最低综合分", "平均综合分", "学考门槛要求"]
C_WIDTHS = [16, 8, 12, 12, 12, 12, 12, 30]
C_NUMFMT = [None, "0", "#,##0", "#,##0", "#,##0", "0.0", "0.0", None]
C_ROWS = []
for s in SCHOOLS:
    for a in s.get("admission", []):
        C_ROWS.append([
            s.get("name"), a.get("year"), a.get("applicants"), a.get("passed"), a.get("admitted"),
            a.get("minScore"), a.get("avgScore"), a.get("xuekaoRequirement"),
        ])

build_sheet(
    wb, "表C-历年竞争比", "表C：历年竞争比（每年 1 行）",
    C_HEADERS, C_ROWS, C_WIDTHS, num_formats=C_NUMFMT,
    center_cols={2}, right_cols={3, 4, 5, 6, 7}, wrap_cols={8}, freeze="B3",
)

# =========================================================
# 表D：报名流程（每环节 1 行）
# =========================================================
D_HEADERS = ["院校全称", "步骤序号", "环节名称", "说明", "截止时间", "需要材料"]
D_WIDTHS = [16, 10, 16, 40, 14, 30]
D_NUMFMT = [None, "0", None, None, None, None]
D_ROWS = []
for s in SCHOOLS:
    for st in sorted(s.get("applicationSteps", []), key=lambda x: x.get("step", 0)):
        D_ROWS.append([
            s.get("name"), st.get("step"), st.get("title"), st.get("description"),
            st.get("deadline"), joinv(st.get("materials")),
        ])

build_sheet(
    wb, "表D-报名流程", "表D：报名流程（每环节 1 行）",
    D_HEADERS, D_ROWS, D_WIDTHS, num_formats=D_NUMFMT,
    center_cols={2}, wrap_cols={4, 6}, freeze="B3",
)

# =========================================================
# 院校清单（2026 全名单 + 历史备查）
# =========================================================
LIST_HEADERS = ["序号", "类别", "院校全称", "id", "状态", "负责人", "备注"]
LIST_WIDTHS = [8, 22, 26, 12, 12, 12, 34]
LIST_ROWS = []


def add_list(cat, pairs, note_fn):
    for name, sid in pairs:
        status = "已收录" if sid in COLLECTED else "待补"
        LIST_ROWS.append([None, cat, name, sid, status, None, note_fn(name, sid, status)])


def _note_difang(n, s, st):
    return "已入库 live-app" if st == "已收录" else None


def _note_zh(n, s, st):
    return "综合评价（非三位一体口径，可按需纳入）"


def _note_ls(n, s, st):
    return "2025 年省内地方属，2026 年名单中已移除"


add_list(CAT_DIFANG, DIFANG_2026, _note_difang)
add_list(CAT_GSP, GSP_2026, _note_difang)
add_list(CAT_ZH, ZH_2026, _note_zh)
add_list(CAT_LS, LS_2026, _note_ls)
for row in LIST_ROWS:
    if row[1] == CAT_LS:
        row[4] = "2026未招"
for i, row in enumerate(LIST_ROWS, start=1):
    row[0] = i

build_sheet(
    wb, "院校清单",
    "浙江省 2026 三位一体 / 综合评价院校全名单 · 省内地方属 32 + 高水平 6 + 综合评价 8 = 46 所（另 2026 未招 7 所备查）",
    LIST_HEADERS, LIST_ROWS, LIST_WIDTHS,
    center_cols={1, 4, 5}, wrap_cols={7}, freeze="A3",
    dv_list=[(2, "省内地方属三位一体,高水平三位一体,综合评价（在浙）,历史（2026未招）"),
             (5, "已收录,待补,已补,已核对,2026未招")],
)

# =========================================================
# 使用说明
# =========================================================
ws_note = wb.create_sheet("使用说明")
ws_note.merge_cells("A1:B1")
t = ws_note["A1"]
t.value = "三位一体数据采集表 · 使用说明"
t.font = font_title
t.fill = FILL_TITLE
t.alignment = Alignment(horizontal="left", vertical="center")
ws_note.row_dimensions[1].height = 28
ws_note.column_dimensions["A"].width = 22
ws_note.column_dimensions["B"].width = 86

NOTE_ROWS = [
    ("【基本信息】", ""),
    ("文档用途", "整理浙江省三位一体 / 综合评价全部院校信息，用于完善系统数据库。"),
    ("数据落点", "live-app/index.html 的 SCHOOLS 数组（live-app/data.js、open-data/schools.json 为其副本）。"),
    ("依据", "live-app/src/types/index.ts 的 School 接口 + index.html 实际渲染字段。"),
    ("更新日期", "2026-09-10"),
    ("本表结构", "表A 院校基础信息／表B 招生专业／表C 历年竞争比／表D 报名流程；四表以「院校全称」关联。另附「院校清单」全名单与待补进度。"),
    ("", ""),
    ("【2026 院校数量（重要）】", ""),
    ("省内地方属三位一体", "32 所（2025 年为 39 所，2026 年减少 7 所）。"),
    ("高水平三位一体", "6 所：复旦大学、上海交通大学、浙江大学、中国科学院大学、西湖大学、宁波东方理工大学（后者 2026 年首次招生）。"),
    ("综合评价（在浙招生）", "8 所：上海纽约大学、昆山杜克大学、北京外国语大学、上海科技大学、南方科技大学、华南理工大学、香港中文大学（深圳）、深圳北理莫斯科大学。"),
    ("2026 合计", "三位一体口径 = 32 + 6 = 38 所；含在浙综合评价共 46 所。"),
    ("当前已收录", "15 所（浙大 + 14 所省内地方属），待补 31 所，详见「院校清单」。"),
    ("2026 未招（备查）", "台州学院、衢州学院、宁波工程学院、温州理工学院、金华职业技术大学、宁波幼儿师范高等专科学校、湖州学院（2025 年曾招生）。"),
    ("", ""),
    ("【格式约定与校验规则】", ""),
    ("权重", "写小数且三项之和 = 1（如 0.15 / 0.25 / 0.60），不要写 15%。"),
    ("学考分值", "学考满分须与 A/B/C/D 自洽（如 A15 B10 C5 D0 → 满分 150）；E 等第不计。"),
    ("最低综合分", "填百分制（0–100，可 1 位小数），不是高考分、不是 750 制。"),
    ("人数 / 计划数", "纯数字，不写「约」「+」「余名」。"),
    ("年份", "4 位数字，如 2026。"),
    ("多值字段", "别名、选科要求、笔试科目、亮点统一用 ; 分隔。"),
    ("缺失值", "留空，不要写「暂无」「-」「待定」，程序按空处理。"),
    ("布尔字段", "统一填「是 / 否」。"),
    ("来源", "涉及数字的字段尽量附来源链接。"),
    ("专业 id", "规则 {schoolId}-m{序号}，如 hdu-m1；「院校清单」中待补院校的 id 为拟定建议，正式收录前须确认（id 一旦使用不可改）。"),
    ("", ""),
    ("【采集优先级】", ""),
    ("第一批 · 计算器可用", "学考 A/B/C/D 分值 + 满分、三项权重、校测满分、高考满分、历年最低综合分、学考门槛。"),
    ("第二批 · 决策信息", "基本信息卡（校区/学费/电话/体检限制）、校测形式与内容、转专业限制、招生专业与选科、报名流程与材料。"),
    ("第三批 · 体验数据", "满意度、宿舍条件、专业录取线、招生简章链接。"),
    ("", ""),
    ("【数据来源与合规】", ""),
    ("2026 三位一体名单", "浙江省教育考试院《2026 三位一体招生开启！报考要点速览》，2026-03-05，https://www.zjzs.net/art/2026/3/5/art_30_12088.html"),
    ("2026 高水平三位一体", "浙江省教育考试院官方微信 6 校简章汇总（复旦/上交/浙大/国科大/西湖/宁波东方理工）。"),
    ("2026 在浙综合评价", "各校 2026 年综合评价招生简章（8 所）。"),
    ("合规提示", "招生政策每年可能调整，务必以官方最新公告为准；涉及人数、分数、学费的数字尽量保留来源链接。"),
    ("数据纠错", "系统内已有「数据纠错」入口，欢迎用户反馈错误数据。"),
]

r = 3
for k, v in NOTE_ROWS:
    if k.startswith("【"):
        ws_note.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        c = ws_note.cell(row=r, column=1, value=k)
        c.font = Font(name=FONT, size=11, bold=True, color=XL_TITLE_BG)
        c.alignment = Alignment(horizontal="left", vertical="center")
    elif k:
        c1 = ws_note.cell(row=r, column=1, value=k)
        c1.font = Font(name=FONT, size=10, bold=True)
        c1.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        c1.fill = FILL_LIGHT
        c1.border = BORDER
        c2 = ws_note.cell(row=r, column=2, value=v)
        c2.font = font_body
        c2.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        c2.border = BORDER
    r += 1

ws_note.freeze_panes = "A3"

wb.properties.title = "三位一体数据采集表"
wb.save(OUT)
print("SAVED:", OUT)
print("sheets:", wb.sheetnames)
print("表A:", len(A_ROWS), "+reserve31 | 表B:", len(B_ROWS), "| 表C:", len(C_ROWS), "| 表D:", len(D_ROWS))
print("院校清单:", len(LIST_ROWS),
      "(地方属", sum(1 for x in LIST_ROWS if x[1] == CAT_DIFANG),
      "/ 高水平", sum(1 for x in LIST_ROWS if x[1] == CAT_GSP),
      "/ 综合评价", sum(1 for x in LIST_ROWS if x[1] == CAT_ZH),
      "/ 历史", sum(1 for x in LIST_ROWS if x[1] == CAT_LS), ")",
      "已收录:", sum(1 for x in LIST_ROWS if x[4] == "已收录"),
      "待补:", sum(1 for x in LIST_ROWS if x[4] == "待补"))
