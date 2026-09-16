# -*- coding: utf-8 -*-
"""生成「三位一体数据采集表-2025.xlsx」。
2025 名单 = 39 所省内地方属 + 5 所高水平 = 44 所（只收录三位一体，排除综合评价 8 所）。
年度字段取自 y2025_1~7.json；跨年稳定字段（校区/学费/满意度/宿舍）复用 2026 数据集 schools_38.json。
"""
import json
import os

try:
    import openpyxl  # noqa
except ImportError:
    import subprocess
    import sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "openpyxl>=3.1.0"])

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

BASE = r"D:\Desktop\VibeCoding\三位一体辅助系统"
REF = os.path.join(BASE, ".三位一体数据采集表.ref")
OUT = os.path.join(BASE, "三位一体数据采集表-2025.xlsx")
OUT_MERGED = os.path.join(REF, "schools_44.json")


def load(fn):
    with open(os.path.join(REF, fn), "r", encoding="utf-8") as fh:
        return json.load(fh)


BASE26 = {s["id"]: s for s in load("schools_38.json")["schools"]}
Y = {}
for i in range(1, 8):
    for s in load("y2025_%d.json" % i):
        Y[s["id"]] = s

ANNUAL = ["formula", "formulaNote", "formulaSource", "examFormat", "transferRestriction",
          "majors", "admission", "applicationSteps", "brochureUrl"]
IDENTITY = ["name", "shortName", "aliases", "type", "info", "satisfaction", "dormitory"]

CAT_DIFANG = "省内地方属三位一体"
CAT_GSP = "高水平三位一体"
CAT_ZH = "综合评价（在浙）"

DIFANG_2025 = [
    ("浙江工业大学", "zjut"), ("浙江师范大学", "zjnu"), ("宁波大学", "nbu"), ("杭州电子科技大学", "hdu"),
    ("浙江工商大学", "zjgsu"), ("浙江理工大学", "zstu"), ("温州医科大学", "wmu"), ("浙江海洋大学", "zjou"),
    ("浙江农林大学", "zafu"), ("浙江中医药大学", "zcmu"), ("中国计量大学", "cjlu"), ("浙江万里学院", "zwu"),
    ("浙江科技大学", "zust"), ("浙江财经大学", "zufe"), ("嘉兴大学", "zjxu"), ("浙大城市学院", "zucc"),
    ("浙大宁波理工学院", "nit"), ("杭州师范大学", "hznu"), ("湖州师范学院", "hzsf"), ("绍兴文理学院", "sxu"),
    ("台州学院", "tzc"), ("温州大学", "wzu"), ("浙江外国语学院", "zisu"), ("宁波工程学院", "nbut"),
    ("衢州学院", "qzc"), ("浙江水利水电学院", "zjweu"), ("浙江警察学院", "zjpc"), ("杭州医学院", "hzmc"),
    ("丽水学院", "lsu"), ("湖州学院", "hzxy"), ("温州理工学院", "wzut"), ("嘉兴南湖学院", "jxnh"),
    ("温州肯恩大学", "wku"), ("宁波诺丁汉大学", "unnc"), ("浙江越秀外国语学院", "zyu"), ("宁波财经学院", "nbfe"),
    ("温州商学院", "wzbc"), ("金华职业技术大学", "jhc"), ("宁波幼儿师范高等专科学校", "nbyz"),
]
GSP_2025 = [("复旦大学", "fudan"), ("上海交通大学", "sjtu"), ("浙江大学", "zju"),
            ("中国科学院大学", "ucas"), ("西湖大学", "westlake")]
ZH_2025 = [("上海纽约大学", "nyush"), ("昆山杜克大学", "dku"), ("北京外国语大学", "bfsu"),
           ("上海科技大学", "shtech"), ("南方科技大学", "sustc"), ("华南理工大学", "scut"),
           ("香港中文大学（深圳）", "cuhksz"), ("深圳北理莫斯科大学", "smbu")]

# 2026-09-16：只收录「三位一体」，排除「综合评价招生」8 所
ORDER = [(n, i, CAT_DIFANG) for n, i in DIFANG_2025] + \
        [(n, i, CAT_GSP) for n, i in GSP_2025]
EXCLUDED_ZH = list(ZH_2025)

# 2026 变化备注
STOPPED_2026 = {"tzc", "qzc", "nbut", "wzut", "jhc", "nbyz", "hzxy"}
RENAMED = {"hzsf": "2026年更名为湖州师范大学", "sxu": "2026年更名为绍兴大学"}

# ---------- 合并 ----------
merged = []
for name, sid, cat in ORDER:
    b = BASE26.get(sid, {})
    yy = Y.get(sid, {})
    s = dict(b)
    for k in IDENTITY:
        if yy.get(k):
            s[k] = yy[k]
    for k in ANNUAL:
        if yy.get(k) is not None:
            s[k] = yy[k]
    if not yy.get("admission"):
        s["admission"] = [a for a in b.get("admission", []) if (a.get("year") or 0) <= 2025]
    s.setdefault("id", sid)
    s["name"] = s.get("name") or name
    merged.append(s)

with open(OUT_MERGED, "w", encoding="utf-8") as fh:
    json.dump({"version": "2025.1", "updatedAt": "2025年（2026-09整理）", "count": len(merged), "schools": merged},
              fh, ensure_ascii=False, indent=2)

BY = {s["id"]: s for s in merged}

# ---------- 样式 ----------
def xl_color(css):
    v = css.removeprefix("#").upper()
    if len(v) != 6:
        raise ValueError(css)
    return "FF" + v


XL_PRIMARY = xl_color("#4472C4")
XL_TITLE_BG = xl_color("#2F5597")
XL_LIGHT = xl_color("#D9E2F3")
XL_BORDER = xl_color("#BFBFBF")
XL_WHITE = xl_color("#FFFFFF")
thin = Side(style="thin", color=XL_BORDER)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
FONT = "微软雅黑"
font_body = Font(name=FONT, size=10)
font_hdr = Font(name=FONT, size=10, bold=True, color=XL_WHITE)
font_title = Font(name=FONT, size=13, bold=True, color=XL_WHITE)
A_HDR = Alignment(horizontal="center", vertical="center", wrap_text=True)
A_LEFT = Alignment(horizontal="left", vertical="center")
A_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)
A_CENTER = Alignment(horizontal="center", vertical="center")
A_RIGHT = Alignment(horizontal="right", vertical="center")
FILL_HDR = PatternFill("solid", fgColor=XL_PRIMARY)
FILL_TITLE = PatternFill("solid", fgColor=XL_TITLE_BG)
FILL_LIGHT = PatternFill("solid", fgColor=XL_LIGHT)
TYPE_MAP = {"ministry": "部属", "provincial": "省属"}
INTERVIEW_MAP = {"individual": "个面", "group": "群面", "both": "两者", "none": None}


def yesno(v):
    return None if v is None else ("是" if v else "否")


def joinv(v):
    return None if not v else ";".join(str(x) for x in v)


def build_sheet(wb, sname, title, headers, rows, widths, num_formats=None,
                center_cols=(), right_cols=(), wrap_cols=(), dv_list=(),
                freeze="B3", use_filter=True):
    ws = wb.create_sheet(sname)
    ncol = len(headers)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    tc = ws.cell(row=1, column=1, value=title)
    tc.font = font_title
    tc.fill = FILL_TITLE
    tc.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28
    for j, h in enumerate(headers, start=1):
        c = ws.cell(row=2, column=j, value=h)
        c.font = font_hdr
        c.fill = FILL_HDR
        c.alignment = A_HDR
        c.border = BORDER
    ws.row_dimensions[2].height = 32
    last = 2 + len(rows)
    for i, rv in enumerate(rows):
        r = 3 + i
        for j in range(1, ncol + 1):
            v = rv[j - 1] if j - 1 < len(rv) else None
            c = ws.cell(row=r, column=j, value=v)
            c.font = font_body
            if j in center_cols:
                c.alignment = A_CENTER
            elif j in right_cols:
                c.alignment = A_RIGHT
            elif j in wrap_cols:
                c.alignment = A_WRAP
            else:
                c.alignment = A_LEFT
            c.border = BORDER
            if num_formats and num_formats[j - 1]:
                c.number_format = num_formats[j - 1]
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    for col_idx, opts in dv_list:
        dv = DataValidation(type="list", formula1='"%s"' % opts, allow_blank=True)
        dv.error = "请从下拉列表中选择"
        ws.add_data_validation(dv)
        col = get_column_letter(col_idx)
        dv.add("%s3:%s%d" % (col, col, last))
    if freeze:
        ws.freeze_panes = freeze
    if use_filter:
        ws.auto_filter.ref = "A2:%s%d" % (get_column_letter(ncol), last)
    return ws


wb = Workbook()
wb.remove(wb.active)

# ===== 表A =====
A_HEADERS = [
    "类别", "院校全称", "简称", "别名", "类型", "id",
    "校区1名称", "校区1地址", "校区2名称", "校区2地址", "官网", "招生办电话", "咨询QQ",
    "普通学费", "中外合作学费", "体检限制",
    "招生简章链接", "章程年份", "章程链接",
    "学考A", "学考B", "学考C", "学考D", "学考满分", "校测满分", "高考满分",
    "学考权重", "校测权重", "高考权重",
    "有笔试", "有面试", "有体测", "笔试科目", "面试形式", "校测内容摘要", "备考建议",
    "转专业受限", "转专业说明",
    "综合满意度", "环境满意度", "生活满意度", "满意度来源",
    "宿舍描述", "宿舍评分", "宿舍亮点", "宿舍不足",
    "折算/口径备注",
]
A_WIDTHS = [18, 14, 9, 14, 8, 9, 12, 16, 12, 16, 26, 14, 12, 18, 20, 24, 26, 10, 26,
            8, 8, 8, 8, 9, 9, 9, 9, 9, 9, 8, 8, 8, 16, 10, 40, 34, 10, 34,
            10, 10, 10, 12, 40, 10, 24, 24, 46]
A_NUMFMT = [None, None, None, None, None, None, None, None, None, None, None, None, None,
            None, None, None, None, "0", None,
            "0", "0", "0", "0", "0", "0", "0", "0.00", "0.00", "0.00",
            None, None, None, None, None, None, None, None, None,
            "0.0", "0.0", "0.0", None, None, "0.0", None, None, None]
A_ROWS = []
for name, sid, cat in ORDER:
    s = BY.get(sid, {})
    info = s.get("info", {}) or {}
    camps = info.get("campuses", []) or []
    c1 = camps[0] if len(camps) > 0 else {}
    c2 = camps[1] if len(camps) > 1 else {}
    f = s.get("formula", {}) or {}
    xk = f.get("xuekao", {}) or {}
    w = f.get("weights", {}) or {}
    ef = s.get("examFormat", {}) or {}
    tr = s.get("transferRestriction", {}) or {}
    sat = s.get("satisfaction", {}) or {}
    dor = s.get("dormitory", {}) or {}
    fs = s.get("formulaSource", {}) or {}
    A_ROWS.append([
        cat, name, s.get("shortName"), joinv(s.get("aliases")),
        TYPE_MAP.get(s.get("type"), s.get("type")), sid,
        c1.get("name"), c1.get("address"), c2.get("name"), c2.get("address"),
        info.get("website"), info.get("admissionsPhone"), info.get("consultQQ"),
        info.get("tuitionGeneral"), info.get("tuitionSinoForeign"), info.get("healthRestrictions"),
        s.get("brochureUrl"), fs.get("year"), fs.get("url"),
        xk.get("A"), xk.get("B"), xk.get("C"), xk.get("D"), xk.get("fullScore"),
        f.get("xiaokao", {}).get("fullScore"), f.get("gaokao", {}).get("fullScore"),
        w.get("xuekao"), w.get("xiaokao"), w.get("gaokao"),
        yesno(ef.get("hasWrittenTest")), yesno(ef.get("hasInterview")), yesno(ef.get("hasPhysicalTest")),
        joinv(ef.get("writtenTestSubjects")),
        INTERVIEW_MAP.get(ef.get("interviewFormat"), ef.get("interviewFormat")),
        ef.get("contentSummary"), ef.get("tips"),
        yesno(tr.get("restricted")), tr.get("detail"),
        sat.get("overall"), sat.get("environment"), sat.get("life"), sat.get("source"),
        dor.get("description"), dor.get("score"), joinv(dor.get("highlights")), joinv(dor.get("drawbacks")),
        s.get("formulaNote"),
    ])

build_sheet(wb, "表A-院校基础信息",
            "表A：院校基础信息（每校 1 行）· 2025 年三位一体 44 所（省属院校 39 + 高水平大学 5）",
            A_HEADERS, A_ROWS, A_WIDTHS, num_formats=A_NUMFMT,
            center_cols={5, 18, 20, 21, 22, 23, 24, 25, 26, 27, 28, 29, 30, 31, 32, 34, 37, 39, 40, 41, 44},
            wrap_cols={8, 10, 13, 16, 19, 33, 35, 36, 38, 43, 45, 46, 47},
            dv_list=[(1, "省内地方属三位一体,高水平三位一体"),
                     (5, "省属,部属"), (30, "是,否"), (31, "是,否"), (32, "是,否"),
                     (34, "个面,群面,两者"), (37, "是,否")],
            freeze="C3")

# ===== 表B =====
B_HEADERS = ["类别", "院校全称", "专业名称", "专业类别", "选科要求", "招生计划数", "最低分年份", "最低综合分", "来源链接"]
B_WIDTHS = [18, 16, 30, 12, 14, 12, 12, 12, 30]
B_NUMFMT = [None, None, None, None, None, "0", "0", "0.0", None]
B_ROWS = []
for name, sid, cat in ORDER:
    for m in BY.get(sid, {}).get("majors", []) or []:
        st = m.get("admissionStats", {}) or {}
        B_ROWS.append([cat, name, m.get("name"), m.get("category"), joinv(m.get("requiredSubjects")),
                       m.get("planCount"), st.get("year"), st.get("minScore"), st.get("sourceUrl")])
build_sheet(wb, "表B-招生专业", "表B：招生专业（每专业 1 行）· 2025 年",
            B_HEADERS, B_ROWS, B_WIDTHS, num_formats=B_NUMFMT,
            center_cols={6, 7, 8}, wrap_cols={9}, freeze="C3")

# ===== 表C =====
C_HEADERS = ["类别", "院校全称", "年份", "报名人数", "入围人数", "录取人数", "最低综合分", "平均综合分", "学考门槛要求"]
C_WIDTHS = [18, 16, 8, 12, 12, 12, 12, 12, 46]
C_NUMFMT = [None, None, "0", "#,##0", "#,##0", "#,##0", "0.0", "0.0", None]
C_ROWS = []
for name, sid, cat in ORDER:
    for a in BY.get(sid, {}).get("admission", []) or []:
        C_ROWS.append([cat, name, a.get("year"), a.get("applicants"), a.get("passed"),
                       a.get("admitted"), a.get("minScore"), a.get("avgScore"), a.get("xuekaoRequirement")])
build_sheet(wb, "表C-历年竞争比", "表C：历年竞争比（每年 1 行）· 2025 年",
            C_HEADERS, C_ROWS, C_WIDTHS, num_formats=C_NUMFMT,
            center_cols={3}, right_cols={4, 5, 6, 7, 8}, wrap_cols={9}, freeze="C3")

# ===== 表D =====
D_HEADERS = ["类别", "院校全称", "步骤序号", "环节名称", "说明", "截止时间", "需要材料"]
D_WIDTHS = [18, 16, 10, 18, 42, 20, 30]
D_NUMFMT = [None, None, "0", None, None, None, None]
D_ROWS = []
for name, sid, cat in ORDER:
    for st in sorted(BY.get(sid, {}).get("applicationSteps", []) or [], key=lambda x: x.get("step", 0)):
        D_ROWS.append([cat, name, st.get("step"), st.get("title"), st.get("description"),
                       st.get("deadline"), joinv(st.get("materials"))])
build_sheet(wb, "表D-报名流程", "表D：报名流程（每环节 1 行）· 2025 年",
            D_HEADERS, D_ROWS, D_WIDTHS, num_formats=D_NUMFMT,
            center_cols={3}, wrap_cols={5, 6, 7}, freeze="C3")

# ===== 院校清单 =====
L_HEADERS = ["序号", "类别", "院校全称", "id", "状态", "负责人", "备注"]
L_WIDTHS = [8, 22, 26, 12, 12, 12, 40]
L_ROWS = []
for idx, (name, sid, cat) in enumerate(ORDER, start=1):
    if sid in STOPPED_2026:
        note = "2026 年起已停止三位一体招生（2026 名单中移除）"
    elif sid in RENAMED:
        note = RENAMED[sid]
    else:
        note = None
    L_ROWS.append([idx, cat, name, sid, "已收录", None, note])
build_sheet(wb, "院校清单",
            "浙江省 2025 三位一体院校全名单 · 39 + 5 = 44 所（均标注 2026 变化）",
            L_HEADERS, L_ROWS, L_WIDTHS, center_cols={1, 4, 5}, wrap_cols={7}, freeze="A3",
            dv_list=[(2, "省内地方属三位一体,高水平三位一体"),
                     (5, "已收录,待补,已补,已核对,已停招")])

# ===== 使用说明 =====
ws_note = wb.create_sheet("使用说明")
ws_note.merge_cells("A1:B1")
t = ws_note["A1"]
t.value = "三位一体数据采集表（2025 年）· 使用说明"
t.font = font_title
t.fill = FILL_TITLE
t.alignment = Alignment(horizontal="left", vertical="center")
ws_note.row_dimensions[1].height = 28
ws_note.column_dimensions["A"].width = 22
ws_note.column_dimensions["B"].width = 96

NOTE_ROWS = [
    ("【基本信息】", ""),
    ("文档用途", "2025 年浙江省「三位一体」院校数据采集表，与 2026 版并存，用于跨年对比与历史查询。**只收录三位一体**。"),
    ("更新日期", "2026-09-10（2026-09-16 调整：仅保留三位一体，剔除综合评价 8 所）"),
    ("本表结构", "表A 院校基础信息／表B 招生专业／表C 历年竞争比／表D 报名流程；含「类别」列与「折算/口径备注」列。"),
    ("", ""),
    ("【2025 三位一体院校数量】", ""),
    ("省属院校三位一体", "39 所（比 2026 年多 7 所）。"),
    ("比 2026 多出的 7 所", "台州学院、衢州学院、宁波工程学院、温州理工学院、金华职业技术大学、宁波幼儿师范高等专科学校、湖州学院（均已于 2026 年停招）。"),
    ("高水平大学三位一体", "5 所：复旦大学、上海交通大学、浙江大学、中国科学院大学、西湖大学（2026 年新增宁波东方理工大学）。"),
    ("合计", "**44 所**（39 + 5）。"),
    ("⚠️ 已排除（非三位一体）", "综合评价招生 8 所（上纽、昆杜、北外、上科大、南科大、华南理工、港中深、深北莫）：官方口径为「综合评价招生」，多数不计学考，与本项目三位一体口径不符，已整体排除。"),
    ("校名变化", "湖州师范学院（2026 年更名为湖州师范大学）、绍兴文理学院（2026 年更名为绍兴大学）。"),
    ("", ""),
    ("【折算口径说明】", ""),
    ("省属院校", "学考 A/B/C/D 分值 + 校测满分 + 高考满分 750，权重三项和为 1。"),
    ("高水平大学", "公式与省属不同（85-10-5、1000分制、60-30-10、70-20-10 等），已按各校 2025 年公布规则填入并在「折算/口径备注」注明，计算器需按类别适配。"),
    ("", ""),
    ("【数据来源与合规】", ""),
    ("名单来源", "浙江省教育考试院《2025 年浙江省普通高校三位一体综合评价招生试点高校名单》(2025-03-01, zjzs.net)。"),
    ("章程来源", "各校本科招生网 2025 年招生章程/简章；高水平取自各校招生网及教育部阳光高考平台。"),
    ("跨年字段", "校区、学费、满意度、宿舍等跨年稳定字段复用 2026 数据集（schools_38.json）。"),
    ("合规提示", "招生政策每年调整，历年数据仅供参考，务必以当年官方最新公告为准。"),
]
r = 3
for k, v in NOTE_ROWS:
    if k.startswith("【"):
        ws_note.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        c = ws_note.cell(row=r, column=1, value=k)
        c.font = Font(name=FONT, size=11, bold=True, color=XL_TITLE_BG)
        c.alignment = Alignment(horizontal="left", vertical="center")
    elif k:
        c1 = ws_note.cell(row=r, column=1, value=k)
        c1.font = Font(name=FONT, size=10, bold=True)
        c1.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        c1.fill = FILL_LIGHT
        c1.border = BORDER
        c2 = ws_note.cell(row=r, column=2, value=v)
        c2.font = font_body
        c2.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        c2.border = BORDER
    r += 1
ws_note.freeze_panes = "A3"

wb.properties.title = "三位一体数据采集表（2025年）"
wb.save(OUT)
print("SAVED:", OUT)
print("MERGED:", OUT_MERGED, len(merged))
print("sheets:", wb.sheetnames)
print("表A:", len(A_ROWS), "| 表B:", len(B_ROWS), "| 表C:", len(C_ROWS), "| 表D:", len(D_ROWS), "| 清单:", len(L_ROWS))
missing = [i for _, i, _ in ORDER if i not in Y]
print("无2025调研数据:", missing)
