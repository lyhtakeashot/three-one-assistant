# -*- coding: utf-8 -*-
"""为 2026 / 2025 两个工作簿各追加一张「折算规则表」：
   一校一行，集中呈现 综合分结算方式（三项比例）+ 学考等级对应分值。
"""
import json
import os

try:
    import openpyxl  # noqa
except ImportError:
    import subprocess
    import sys
    subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet", "openpyxl>=3.1.0"])

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BASE = r"D:\Desktop\VibeCoding\三位一体辅助系统"
REF = os.path.join(BASE, ".三位一体数据采集表.ref")

CATS = ["省内地方属三位一体", "高水平三位一体"]

TARGETS = [
    (os.path.join(BASE, "三位一体数据采集表.xlsx"), "schools_38.json", 2026, (32, 6)),
    (os.path.join(BASE, "三位一体数据采集表-2025.xlsx"), "schools_44.json", 2025, (39, 5)),
]

# ---------- 修正溯源：读 6 批章程核对结果 ----------
VERIFY = []
for _i in range(1, 7):
    _p = os.path.join(REF, "verify_out_%d.json" % _i)
    if os.path.exists(_p):
        with open(_p, "r", encoding="utf-8") as _fh:
            VERIFY += json.load(_fh)
# 只保留「需留档」的条目：已修正 + 无法核实（后者也留档，状态标「待核实」）
APPLIED = []
for e in VERIFY:
    if e.get("ok"):
        continue
    off = e.get("official") or {}
    if not off or all(off.get(k) in (None, {}) for k in ("xuekao", "xiaokaoFS", "weights")):
        if e.get("diff") or e.get("note"):
            e = dict(e)
            e["_unverified"] = True
            APPLIED.append(e)
        continue
    APPLIED.append(e)


def fix_rows_for(year, title_of):
    rows = []
    for i, e in enumerate([x for x in APPLIED if x.get("year") == year], start=1):
        for item in (e.get("diff") or ["（口径/说明性修订）"]):
            if e.get("_unverified") or ("无法核实" in item) or ("未提供" in item) or ("链接失效" in item):
                status = "待核实"
            elif ("未明确" in item) or ("未公布" in item):
                status = "说明性（已补注）"
            else:
                status = "已修正"
            rows.append([
                i, year, e["name"], e["id"], item, status,
                e.get("source"), title_of.get(e["id"], (e["name"] + " 招生章程")),
                e.get("quote") or "", e.get("note") or "", "2026-09-17",
            ])
    return rows


FIX_HEADERS = ["序号", "年份", "院校全称", "id", "修正项（原值 → 官方值）", "状态",
               "来源链接", "来源标题", "章程原文摘录", "备注", "核对日期"]
FIX_WIDTHS = [6, 8, 18, 9, 52, 14, 46, 34, 60, 40, 12]


def xl_color(css):
    v = css.removeprefix("#").upper()
    if len(v) != 6:
        raise ValueError(css)
    return "FF" + v


XL_PRIMARY = xl_color("#4472C4")
XL_TITLE_BG = xl_color("#2F5597")
XL_BORDER = xl_color("#BFBFBF")
XL_WHITE = xl_color("#FFFFFF")
thin = Side(style="thin", color=XL_BORDER)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
FONT = "微软雅黑"
font_body = Font(name=FONT, size=10)
font_hdr = Font(name=FONT, size=10, bold=True, color=XL_WHITE)
font_title = Font(name=FONT, size=13, bold=True, color=XL_WHITE)
A_HDR = Alignment(horizontal="center", vertical="center", wrap_text=True)
A_LEFT = Alignment(horizontal="left", vertical="center")
A_WRAP = Alignment(horizontal="left", vertical="top", wrap_text=True)
A_CENTER = Alignment(horizontal="center", vertical="center")
FILL_HDR = PatternFill("solid", fgColor=XL_PRIMARY)
FILL_TITLE = PatternFill("solid", fgColor=XL_TITLE_BG)

HEADERS = ["序号", "类别", "院校全称", "id",
           "学考A", "学考B", "学考C", "学考D", "学考满分",
           "校测满分", "高考满分",
           "学考比例", "校测比例", "高考比例",
           "综合分结算方式", "折算说明", "章程年份", "章程链接"]
WIDTHS = [6, 18, 16, 9, 8, 8, 8, 8, 9, 9, 9, 9, 9, 9, 62, 44, 9, 34]
NUMFMT = [None, None, None, None, "0.0", "0.0", "0.0", "0.0", "0.0", "0.0", "0",
          "0%", "0%", "0%", None, None, "0", None]
CENTER = {1, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 17}
WRAP = {15, 16, 18}


def pct(v):
    if v is None:
        return None
    x = round(v * 100, 4)
    if abs(x - round(x)) < 1e-9:
        return "%d%%" % round(x)
    return ("%g%%" % x)


def score_txt(label, fs, weight):
    w = pct(weight)
    if fs:
        return "%s(满分%g)×%s" % (label, fs, w)
    return "%s×%s" % (label, w)


def build_rows(schools, counts, year):
    rows = []
    i = 0
    for cat, n in zip(CATS, counts):
        for s in schools[i:i + n]:
            f = s.get("formula") or {}
            xk = f.get("xuekao") or {}
            w = f.get("weights") or {}
            fs = s.get("formulaSource") or {}
            note = s.get("formulaNote") or ""
            extra = [k for k in ("E",) if k in xk]
            if extra:
                seg = "；".join("%s=%g" % (k, xk[k]) for k in extra)
                note = (note + "；" if note else "") + "另有 " + seg
            term = " + ".join([
                score_txt("学考折算分", xk.get("fullScore"), w.get("xuekao")),
                score_txt("综合测试分", (f.get("xiaokao") or {}).get("fullScore"), w.get("xiaokao")),
                score_txt("高考分", (f.get("gaokao") or {}).get("fullScore"), w.get("gaokao")),
            ])
            rows.append([
                None, cat, s.get("name"), s.get("id"),
                xk.get("A"), xk.get("B"), xk.get("C"), xk.get("D"), xk.get("fullScore"),
                (f.get("xiaokao") or {}).get("fullScore"), (f.get("gaokao") or {}).get("fullScore"),
                w.get("xuekao"), w.get("xiaokao"), w.get("gaokao"),
                "综合分 = " + term, note or None,
                fs.get("year") or year, fs.get("url") or s.get("brochureUrl"),
            ])
        i += n
    for k, r in enumerate(rows, start=1):
        r[0] = k
    return rows


for xlsx, jsname, year, counts in TARGETS:
    with open(os.path.join(REF, jsname), "r", encoding="utf-8") as fh:
        schools = json.load(fh)["schools"]
    rows = build_rows(schools, counts, year)

    wb = load_workbook(xlsx)
    if "折算规则表" in wb.sheetnames:
        del wb["折算规则表"]
    idx = wb.sheetnames.index("院校清单") if "院校清单" in wb.sheetnames else len(wb.sheetnames)
    ws = wb.create_sheet("折算规则表", idx)

    ncol = len(HEADERS)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    tc = ws.cell(row=1, column=1,
                 value="分数结算方式与学考等级折算分值 · %d 年（共 %d 所）· 综合分 = 学考折算分×比例 + 综合测试分×比例 + 高考分×比例" % (year, len(rows)))
    tc.font = font_title
    tc.fill = FILL_TITLE
    tc.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28

    for j, h in enumerate(HEADERS, start=1):
        c = ws.cell(row=2, column=j, value=h)
        c.font = font_hdr
        c.fill = FILL_HDR
        c.alignment = A_HDR
        c.border = BORDER
    ws.row_dimensions[2].height = 30

    for i, rv in enumerate(rows):
        r = 3 + i
        for j in range(1, ncol + 1):
            c = ws.cell(row=r, column=j, value=rv[j - 1])
            c.font = font_body
            c.alignment = A_CENTER if j in CENTER else (A_WRAP if j in WRAP else A_LEFT)
            c.border = BORDER
            if NUMFMT[j - 1]:
                c.number_format = NUMFMT[j - 1]

    for j, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w

    last = 2 + len(rows)
    ws.auto_filter.ref = "A2:%s%d" % (get_column_letter(ncol), last)
    ws.freeze_panes = "E3"

    # ---- 修正溯源表：逐项对应官方章程链接与原文 ----
    if "修正溯源" in wb.sheetnames:
        del wb["修正溯源"]
    idx2 = wb.sheetnames.index("院校清单") if "院校清单" in wb.sheetnames else len(wb.sheetnames)
    ws2 = wb.create_sheet("修正溯源", idx2)
    title_of = {s.get("id"): ((s.get("formulaSource") or {}).get("title") or (str(s.get("name")) + " 招生章程"))
                for s in schools}
    frows = fix_rows_for(year, title_of)
    n2 = len(FIX_HEADERS)
    ws2.merge_cells(start_row=1, start_column=1, end_row=1, end_column=n2)
    t2 = ws2.cell(row=1, column=1,
                  value="%d 年折算口径修正溯源 · 每项修正对应官方章程链接与原文摘录（章程核对日期 2026-09-17）" % year)
    t2.font = font_title
    t2.fill = FILL_TITLE
    t2.alignment = Alignment(horizontal="left", vertical="center")
    ws2.row_dimensions[1].height = 28
    for j, h in enumerate(FIX_HEADERS, start=1):
        c = ws2.cell(row=2, column=j, value=h)
        c.font = font_hdr
        c.fill = FILL_HDR
        c.alignment = A_HDR
        c.border = BORDER
    ws2.row_dimensions[2].height = 30
    for i2, rv in enumerate(frows):
        r2 = 3 + i2
        for j in range(1, n2 + 1):
            c = ws2.cell(row=r2, column=j, value=rv[j - 1])
            c.font = font_body
            c.alignment = A_CENTER if j in (1, 2, 4, 6, 11) else A_WRAP
            c.border = BORDER
    for j, w in enumerate(FIX_WIDTHS, start=1):
        ws2.column_dimensions[get_column_letter(j)].width = w
    ws2.freeze_panes = "C3"

    wb.save(xlsx)
    print("OK", os.path.basename(xlsx), "折算规则表", len(rows), "行 | 修正溯源", len(frows), "行 | sheets:", wb.sheetnames)
